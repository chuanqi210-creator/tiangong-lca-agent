"""Deterministic, local-only processors for ordinary research documents."""

from __future__ import annotations

import codecs
import csv
import hashlib
import io
import json
import re
import struct
import wave
import zipfile
from collections.abc import Callable, Mapping
from copy import deepcopy
from html.parser import HTMLParser
from typing import Any, BinaryIO
from xml.etree import ElementTree


MAX_INPUT_BYTES = 100 * 1024 * 1024
MAX_IMAGE_PIXELS = 40_000_000
MAX_PDF_PAGES = 500
MAX_PDF_ACTION_NODES = 10_000
MAX_PDF_ACTION_DEPTH = 32
MAX_OOXML_XML_BYTES = 10 * 1024 * 1024
MAX_OOXML_TOTAL_XML_BYTES = 50 * 1024 * 1024
MAX_OOXML_XML_PARTS = 10_000


def _read_bounded(stream: BinaryIO, *, maximum: int = MAX_INPUT_BYTES) -> bytes:
    chunks, total = [], 0
    while chunk := stream.read(min(1024 * 1024, maximum + 1 - total)):
        if not isinstance(chunk, bytes):
            raise ValueError("processor input stream must yield bytes")
        total += len(chunk)
        if total > maximum:
            raise ValueError("processor input exceeds output-safe ceiling")
        chunks.append(chunk)
    return b"".join(chunks)


def _zip_names(payload: bytes) -> set[str]:
    try:
        with zipfile.ZipFile(io.BytesIO(payload)) as archive:
            return set(archive.namelist())
    except (OSError, zipfile.BadZipFile):
        return set()


def detect_document_type(payload: bytes) -> dict[str, Any]:
    """Detect from signatures and parsable structure; filenames are never consulted."""

    signatures: list[str] = []
    mime = "application/octet-stream"
    if payload.startswith(b"%PDF-"):
        return {"mime": "application/pdf", "signatures": ["pdf_header"]}
    if payload.startswith(b"\x89PNG\r\n\x1a\n"):
        return {"mime": "image/png", "signatures": ["png_signature"]}
    if payload.startswith(b"\xff\xd8\xff"):
        return {"mime": "image/jpeg", "signatures": ["jpeg_soi"]}
    if payload[:4] == b"RIFF" and payload[8:12] == b"WAVE":
        return {"mime": "audio/wav", "signatures": ["riff", "wave"]}
    if payload.startswith(b"ID3") or (len(payload) > 1 and payload[0] == 0xFF and payload[1] & 0xE0 == 0xE0):
        return {"mime": "audio/mpeg", "signatures": ["mpeg_audio_frame"]}
    if payload.startswith(b"Rar!\x1a\x07"):
        return {"mime": "application/vnd.rar", "signatures": ["rar_signature"]}
    if payload.startswith(b"7z\xbc\xaf'\x1c"):
        return {"mime": "application/x-7z-compressed", "signatures": ["7z_signature"]}
    if payload.startswith(b"SQLite format 3\x00"):
        return {"mime": "application/x-sqlite3", "signatures": ["sqlite3_header"]}
    if len(payload) >= 12 and payload[4:8] == b"ftyp":
        return {"mime": "audio/mp4", "signatures": ["iso_bmff"]}
    if payload.startswith(b"PK\x03\x04"):
        names = _zip_names(payload)
        if "word/document.xml" in names:
            mime, signatures = "application/vnd.openxmlformats-officedocument.wordprocessingml.document", ["zip", "opc", "word/document.xml"]
        elif "ppt/presentation.xml" in names:
            mime, signatures = "application/vnd.openxmlformats-officedocument.presentationml.presentation", ["zip", "opc", "ppt/presentation.xml"]
        elif "xl/workbook.xml" in names:
            mime, signatures = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", ["zip", "opc", "xl/workbook.xml"]
        else:
            mime, signatures = "application/zip", ["zip"]
        return {"mime": mime, "signatures": signatures}
    if b"\x00" in payload[:4096]:
        return {"mime": mime, "signatures": []}
    try:
        text = payload.decode("utf-8")
    except UnicodeDecodeError:
        return {"mime": mime, "signatures": []}
    stripped = text.lstrip("\ufeff\r\n\t ")
    try:
        json_value = json.loads(text)
    except (json.JSONDecodeError, RecursionError):
        pass
    else:
        if isinstance(json_value, (dict, list)):
            return {"mime": "application/json", "signatures": ["utf8", "strict_json", "structured_root"]}
    if stripped.startswith("<"):
        if re.match(r"(?is)^<(?:!doctype\s+html|html(?:\s|>))", stripped):
            return {"mime": "text/html", "signatures": ["utf8", "html_root"]}
        try:
            if "<!DOCTYPE" not in text.upper() and "<!ENTITY" not in text.upper():
                ElementTree.fromstring(text)
        except ElementTree.ParseError:
            pass
        else:
            return {"mime": "application/xml", "signatures": ["utf8", "xml_root"]}
    lines = [line for line in text.splitlines() if line]
    for delimiter, mime in (("\t", "text/tab-separated-values"), (",", "text/csv")):
        widths = [len(next(csv.reader([line], delimiter=delimiter))) for line in lines]
        if len(widths) >= 2 and widths[0] >= 2 and len(set(widths)) == 1:
            return {"mime": mime, "signatures": ["utf8", "consistent_tabular_shape"]}
    return {"mime": "text/plain", "signatures": ["utf8"]}


Processor = Callable[[Mapping[str, Any], BinaryIO, Mapping[str, Any]], dict[str, Any]]


class ProcessorRegistry:
    def __init__(self) -> None:
        self._processors: dict[str, Processor] = {}

    def register(self, mime: str, processor: Processor) -> None:
        if not isinstance(mime, str) or "/" not in mime or not callable(processor):
            raise ValueError("processor registration is invalid")
        self._processors[mime] = processor

    def process(
        self,
        artifact: Mapping[str, Any],
        payload: bytes,
        detection: Mapping[str, Any],
    ) -> dict[str, Any]:
        processor = self._processors.get(str(detection.get("mime")))
        if processor is None:
            return {
                "processor": {"id": "deterministic-local-unsupported", "version": "1.0.0"},
                "processing_mode": "metadata_only",
                "status": "unsupported",
                "units": [],
                "warnings": ["no deterministic local processor for detected signature"],
                "detected_type": deepcopy(dict(detection)),
            }
        return processor(artifact, io.BytesIO(payload), detection)


def _result(processor_id: str, detection: Mapping[str, Any], *, units: list[dict[str, Any]], status: str = "ready", warnings: list[str] | None = None, mode: str = "structured_document_parse") -> dict[str, Any]:
    for unit in units:
        original = unit.get("content")
        if original is None:
            original = deepcopy(unit.get("cells", unit.get("metadata")))
        unit.setdefault(
            "provenance",
            {
                "original_value": original,
                "original_unit": None,
                "transformations": [],
                "assumptions": [],
                "missing_data": [],
                "validation_status": "candidate_unvalidated",
            },
        )
    return {
        "processor": {"id": processor_id, "version": "1.0.0"},
        "processing_mode": mode,
        "status": status,
        "units": units,
        "warnings": list(warnings or []),
        "detected_type": deepcopy(dict(detection)),
    }


def process_text(_artifact: Mapping[str, Any], stream: BinaryIO, detection: Mapping[str, Any]) -> dict[str, Any]:
    payload = _read_bounded(stream)
    try:
        text = payload.decode("utf-8-sig", errors="strict")
    except UnicodeDecodeError:
        return _result("bounded-utf8", detection, units=[], status="needs_user", warnings=["UTF-8 decoding failed; encoding diagnosis requires user input"])
    units, offset = [], 0
    for line_number, line in enumerate(text.splitlines(keepends=True), 1):
        content = line.rstrip("\r\n")
        if content:
            units.append({"kind": "text", "content": content, "source_span": {"kind": "text_span", "locator": {"start_char": offset, "end_char": offset + len(content), "start_line": line_number, "end_line": line_number}}, "candidate_only": True})
        offset += len(line)
    if text and not units:
        units.append({"kind": "text", "content": text, "source_span": {"kind": "text_span", "locator": {"start_char": 0, "end_char": len(text), "start_line": 1, "end_line": 1}}, "candidate_only": True})
    return _result("bounded-utf8", detection, units=units)


def _pointer_token(value: str) -> str:
    return value.replace("~", "~0").replace("/", "~1")


def process_json(_artifact: Mapping[str, Any], stream: BinaryIO, detection: Mapping[str, Any]) -> dict[str, Any]:
    value = json.loads(_read_bounded(stream).decode("utf-8-sig", errors="strict"))
    units: list[dict[str, Any]] = []
    def visit(item: Any, pointer: str) -> None:
        if isinstance(item, dict):
            for key, child in item.items():
                visit(child, pointer + "/" + _pointer_token(str(key)))
        elif isinstance(item, list):
            for index, child in enumerate(item):
                visit(child, pointer + f"/{index}")
        else:
            units.append({"kind": "json_value", "content": item, "source_span": {"kind": "json_pointer", "locator": {"pointer": pointer}}, "candidate_only": True})
    visit(value, "")
    if not units:
        units.append({"kind": "json_value", "content": value, "source_span": {"kind": "json_pointer", "locator": {"pointer": ""}}, "candidate_only": True})
    return _result("strict-json-pointer", detection, units=units)


def _xml_name(tag: str, namespaces: Mapping[str, str] | None = None) -> str:
    uri, separator, local = tag[1:].partition("}") if tag.startswith("{") else ("", "", tag)
    value = local if separator else tag.split(":", 1)[-1]
    value = re.sub(r"[^A-Za-z0-9_.-]", "_", value)
    value = value if re.match(r"[A-Za-z_]", value) else "node"
    prefix = namespaces.get(uri) if namespaces and uri else None
    return f"{prefix}:{value}" if prefix else value


def process_xml(_artifact: Mapping[str, Any], stream: BinaryIO, detection: Mapping[str, Any]) -> dict[str, Any]:
    payload = _read_bounded(stream)
    upper = payload.upper()
    if b"<!DOCTYPE" in upper or b"<!ENTITY" in upper:
        return _result("safe-xml", detection, units=[], status="unsupported", warnings=["external entity or DTD content is forbidden"])
    root = ElementTree.fromstring(payload)
    namespace_uris: set[str] = set()
    for element in root.iter():
        if element.tag.startswith("{"):
            uri = element.tag[1:].split("}", 1)[0]
            namespace_uris.add(uri)
    namespaces = {
        uri: f"ns{hashlib.sha256(uri.encode()).hexdigest()[:12]}"
        for uri in namespace_uris
    }
    units: list[dict[str, Any]] = []
    def walk(element: ElementTree.Element, xpath: str) -> None:
        content = (element.text or "").strip()
        if content:
            units.append({"kind": "xml_text", "content": content, "source_span": {"kind": "xml_xpath", "locator": {"xpath": xpath}}, "candidate_only": True})
        counts: dict[str, int] = {}
        for child in list(element):
            name = _xml_name(child.tag, namespaces); counts[name] = counts.get(name, 0) + 1
            walk(child, f"{xpath}/{name}[{counts[name]}]")
    walk(root, "/" + _xml_name(root.tag, namespaces))
    return _result("safe-xml", detection, units=units)


class _SafeHTML(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.excluded_depth = 0; self.block: str | None = None; self.text: list[str] = []; self.sections: list[tuple[str, str, str]] = []
    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag in {"script", "style", "iframe", "object", "embed"}:
            self.excluded_depth += 1
        elif not self.excluded_depth and tag in {"h1", "h2", "h3", "h4", "h5", "h6", "p", "li", "td", "th"}:
            self.block = tag; self.text = []
    def handle_endtag(self, tag: str) -> None:
        if tag in {"script", "style", "iframe", "object", "embed"} and self.excluded_depth:
            self.excluded_depth -= 1
        elif not self.excluded_depth and self.block == tag:
            original = "".join(self.text)
            value = " ".join(original.split())
            if value: self.sections.append((tag, value, original))
            self.block = None; self.text = []
    def handle_data(self, data: str) -> None:
        if not self.excluded_depth and self.block:
            self.text.append(data)


def process_html(_artifact: Mapping[str, Any], stream: BinaryIO, detection: Mapping[str, Any]) -> dict[str, Any]:
    text = _read_bounded(stream).decode("utf-8-sig", errors="strict")
    parser = _SafeHTML(); parser.feed(text); parser.close()
    units, heading, paragraph = [], "Document", 0
    for tag, content, original in parser.sections:
        if tag.startswith("h"):
            heading = content; paragraph = 0
        else:
            units.append({"kind": "document_text", "content": content, "provenance": {"original_value": original, "original_unit": None, "transformations": ["normalized_whitespace"] if original != content else [], "assumptions": [], "missing_data": [], "validation_status": "candidate_unvalidated"}, "source_span": {"kind": "document_section", "locator": {"heading_path": [heading], "paragraph_start": paragraph, "paragraph_end": paragraph}}, "candidate_only": True})
            paragraph += 1
    warnings = ["active content excluded and never executed"] if re.search(r"(?is)<(?:script|style|iframe|object|embed)\b", text) else []
    return _result("inert-html", detection, units=units, warnings=warnings)


def _unit_like_headers(values: list[str]) -> list[str]:
    units = []
    for value in values:
        match = re.search(r"(?:\(|\[|/|\b)(kg|g|mg|t|kwh|mj|m3|l|mol|%)\b", value.casefold())
        if match and match.group(1) not in units: units.append(match.group(1))
    return units


def _table_units(rows: list[list[Any]], *, sheet_name: str) -> list[dict[str, Any]]:
    from openpyxl.utils import get_column_letter
    headers = ["" if value is None else str(value) for value in (rows[0] if rows else [])]
    units = []
    for index, row in enumerate(rows, 1):
        cells = []
        for value in row:
            formula = isinstance(value, str) and value.startswith("=")
            cells.append({"raw_value": value, "displayed_value": value, "formula": formula, "calculated_truth": False})
        end_column = max(1, len(cells))
        units.append({"kind": "table_row", "cells": cells, "unit_like_headers": _unit_like_headers(headers), "source_span": {"kind": "spreadsheet_range", "locator": {"sheet_name": sheet_name, "cell_range": f"A{index}:{get_column_letter(end_column)}{index}"}}, "candidate_only": True})
    return units


def process_delimited(_artifact: Mapping[str, Any], stream: BinaryIO, detection: Mapping[str, Any]) -> dict[str, Any]:
    text = _read_bounded(stream).decode("utf-8-sig", errors="strict")
    delimiter = "\t" if detection["mime"] == "text/tab-separated-values" else ","
    rows = list(csv.reader(io.StringIO(text), delimiter=delimiter))
    return _result("strict-delimited-table", detection, units=_table_units(rows, sheet_name="Table1"), mode="tabular_parse")


def _decode_xml_part(content: bytes) -> str:
    if content.startswith(codecs.BOM_UTF32_LE):
        encoding = "utf-32-le"
    elif content.startswith(codecs.BOM_UTF32_BE):
        encoding = "utf-32-be"
    elif content.startswith(codecs.BOM_UTF16_LE):
        encoding = "utf-16-le"
    elif content.startswith(codecs.BOM_UTF16_BE):
        encoding = "utf-16-be"
    elif content.startswith(codecs.BOM_UTF8):
        encoding = "utf-8-sig"
    elif content.startswith(b"\x00\x00\x00<"):
        encoding = "utf-32-be"
    elif content.startswith(b"<\x00\x00\x00"):
        encoding = "utf-32-le"
    elif content.startswith(b"\x00<\x00?"):
        encoding = "utf-16-be"
    elif content.startswith(b"<\x00?\x00"):
        encoding = "utf-16-le"
    else:
        declaration = content[:256].decode("ascii", errors="ignore")
        match = re.match(
            r'''(?is)^\s*<\?xml\b[^>]*\bencoding\s*=\s*["']([a-z0-9._-]+)["']''',
            declaration,
        )
        encoding = (match.group(1).casefold() if match else "utf-8")
        allowed = {
            "utf-8": "utf-8",
            "us-ascii": "ascii",
            "iso-8859-1": "iso-8859-1",
            "windows-1252": "windows-1252",
        }
        if encoding not in allowed:
            raise UnicodeError("unsupported OOXML XML encoding")
        encoding = allowed[encoding]
    return content.decode(encoding, errors="strict")


def _office_active_warnings(payload: bytes) -> list[str]:
    warnings: list[str] = []
    unsafe_xml: list[str] = []
    macro_content_type = False
    external_relationship = False
    try:
        with zipfile.ZipFile(io.BytesIO(payload)) as archive:
            names = {name.casefold() for name in archive.namelist()}
            xml_members = [
                info for info in archive.infolist()
                if info.filename.casefold().endswith((".xml", ".rels"))
            ]
            if len(xml_members) > MAX_OOXML_XML_PARTS:
                unsafe_xml.append("OOXML XML part-count ceiling")
                xml_members = xml_members[:MAX_OOXML_XML_PARTS]
            total = 0
            for info in xml_members:
                total += info.file_size
                if (
                    info.file_size > MAX_OOXML_XML_BYTES
                    or total > MAX_OOXML_TOTAL_XML_BYTES
                ):
                    unsafe_xml.append(info.filename)
                    continue
                with archive.open(info) as member:
                    content = member.read(MAX_OOXML_XML_BYTES + 1)
                if len(content) > MAX_OOXML_XML_BYTES:
                    unsafe_xml.append(info.filename)
                    continue
                try:
                    decoded = _decode_xml_part(content)
                except UnicodeError:
                    unsafe_xml.append(info.filename)
                    continue
                upper = decoded.upper()
                if "<!DOCTYPE" in upper or "<!ENTITY" in upper:
                    unsafe_xml.append(info.filename)
                if info.filename.casefold() == "[content_types].xml" and (
                    "vnd.ms-office.vbaproject" in decoded.casefold()
                ):
                    macro_content_type = True
                if info.filename.casefold().endswith(".rels") and re.search(
                    r'''(?i)TargetMode\s*=\s*["']External["']''', decoded
                ):
                    external_relationship = True
    except (OSError, zipfile.BadZipFile):
        return ["invalid OOXML package"]
    if any(name.endswith("vbaproject.bin") for name in names) or macro_content_type:
        warnings.append("macro project reported and never executed")
    if unsafe_xml:
        warnings.append("DTD or entity declaration reported and never resolved")
    if external_relationship or any(
        name.endswith("externallinks.xml") or "/externallinks/" in name
        for name in names
    ):
        warnings.append("external relationship reported and never resolved")
    return warnings


_PDF_ACTION_SUBTYPES = {
    "GoTo", "GoToR", "GoToE", "Launch", "Thread", "URI", "Sound",
    "Movie", "Hide", "Named", "SubmitForm", "ResetForm", "ImportData",
    "JavaScript", "SetOCGState", "Rendition", "Trans", "GoTo3DView",
    "RichMediaExecute",
}


def _inspect_pdf_actions(reader: Any) -> dict[str, Any]:
    from pypdf.generic import ArrayObject, DictionaryObject, IndirectObject

    found: set[str] = set()
    seen: set[tuple[int, int] | int] = set()
    nodes = 0
    incomplete: set[str] = set()

    def resolve(value: Any, depth: int) -> Any | None:
        nonlocal nodes
        if depth > MAX_PDF_ACTION_DEPTH:
            incomplete.add("depth ceiling reached")
            return None
        if nodes >= MAX_PDF_ACTION_NODES:
            incomplete.add("node ceiling reached")
            return None
        if isinstance(value, IndirectObject):
            identity = (value.idnum, value.generation)
            if identity in seen:
                return None
            seen.add(identity)
            try:
                value = value.get_object()
            except Exception:
                incomplete.add("indirect object resolution failed")
                return None
        elif isinstance(value, (DictionaryObject, ArrayObject)):
            identity = id(value)
            if identity in seen:
                return None
            seen.add(identity)
        else:
            return value
        nodes += 1
        return value

    def actions(value: Any, depth: int) -> None:
        resolved = resolve(value, depth)
        if isinstance(resolved, ArrayObject):
            for child in resolved:
                actions(child, depth + 1)
            return
        if not isinstance(resolved, DictionaryObject):
            incomplete.add("invalid action entry")
            return
        subtype = resolved.get("/S")
        name = subtype[1:] if isinstance(subtype, str) and subtype.startswith("/") else ""
        if name not in _PDF_ACTION_SUBTYPES:
            incomplete.add("invalid action dictionary")
            return
        found.add(name)
        if "/Next" in resolved:
            actions(resolved["/Next"], depth + 1)

    def additional_actions(value: Any, depth: int) -> None:
        resolved = resolve(value, depth)
        if not isinstance(resolved, DictionaryObject):
            incomplete.add("invalid additional-action dictionary")
            return
        for action in resolved.values():
            actions(action, depth + 1)

    def action_entries(value: Any, depth: int) -> DictionaryObject | None:
        resolved = resolve(value, depth)
        if not isinstance(resolved, DictionaryObject):
            incomplete.add("invalid PDF action owner")
            return None
        if "/A" in resolved:
            actions(resolved["/A"], depth + 1)
        if "/AA" in resolved:
            additional_actions(resolved["/AA"], depth + 1)
        return resolved

    def annotations(value: Any, depth: int) -> None:
        resolved = resolve(value, depth)
        if not isinstance(resolved, ArrayObject):
            incomplete.add("invalid annotation array")
            return
        for annotation in resolved:
            action_entries(annotation, depth + 1)

    def outline(value: Any, depth: int) -> None:
        item = action_entries(value, depth)
        if item is None:
            return
        if "/First" in item:
            outline(item["/First"], depth + 1)
        if "/Next" in item:
            outline(item["/Next"], depth + 1)

    def fields(value: Any, depth: int) -> None:
        resolved = resolve(value, depth)
        if not isinstance(resolved, ArrayObject):
            incomplete.add("invalid form-field array")
            return
        for field_value in resolved:
            field = action_entries(field_value, depth + 1)
            if field is not None and "/Kids" in field:
                fields(field["/Kids"], depth + 1)

    def presentation_steps(value: Any, depth: int) -> None:
        node = resolve(value, depth)
        if node is None:
            return
        if not isinstance(node, DictionaryObject):
            incomplete.add("invalid presentation-step navigation node")
            return
        node_type = node.get("/Type")
        if node_type is not None and node_type != "/NavNode":
            incomplete.add("invalid presentation-step navigation-node type")
        for key in ("/NA", "/PA"):
            if key in node:
                actions(node[key], depth + 1)
        for key in ("/Next", "/Prev"):
            if key in node:
                target = node[key]
                if not isinstance(target, (DictionaryObject, IndirectObject)):
                    incomplete.add("invalid presentation-step navigation link")
                    continue
                presentation_steps(target, depth + 1)

    def javascript_name_tree(value: Any, depth: int) -> None:
        tree = resolve(value, depth)
        if not isinstance(tree, DictionaryObject):
            incomplete.add("invalid JavaScript name tree")
            return
        if "/Names" in tree:
            names = resolve(tree["/Names"], depth + 1)
            if not isinstance(names, ArrayObject) or len(names) % 2:
                incomplete.add("invalid JavaScript name array")
            else:
                for index in range(1, len(names), 2):
                    actions(names[index], depth + 2)
        if "/Kids" in tree:
            kids = resolve(tree["/Kids"], depth + 1)
            if not isinstance(kids, ArrayObject):
                incomplete.add("invalid JavaScript name-tree children")
            else:
                for child in kids:
                    javascript_name_tree(child, depth + 2)

    root = resolve(reader.trailer.get("/Root"), 0)
    if not isinstance(root, DictionaryObject):
        incomplete.add("invalid PDF catalog")
        return {"subtypes": [], "complete": False, "reasons": sorted(incomplete)}
    if "/OpenAction" in root:
        open_action = root["/OpenAction"]
        if isinstance(open_action, (DictionaryObject, IndirectObject)):
            actions(open_action, 1)
        elif not isinstance(open_action, (ArrayObject, str)):
            incomplete.add("invalid catalog open action")
    if "/AA" in root:
        additional_actions(root["/AA"], 1)
    if "/Outlines" in root:
        outlines = resolve(root["/Outlines"], 1)
        if isinstance(outlines, DictionaryObject) and "/First" in outlines:
            outline(outlines["/First"], 2)
    if "/Names" in root:
        names = resolve(root["/Names"], 1)
        if isinstance(names, DictionaryObject) and "/JavaScript" in names:
            javascript_name_tree(names["/JavaScript"], 2)
    if "/AcroForm" in root:
        form = action_entries(root["/AcroForm"], 1)
        if form is not None and "/Fields" in form:
            fields(form["/Fields"], 2)
    for page in reader.pages:
        page_object = action_entries(page, 1)
        if page_object is not None:
            if "/Annots" in page_object:
                annotations(page_object["/Annots"], 2)
            if "/PresSteps" in page_object:
                presentation_steps(page_object["/PresSteps"], 2)
    return {
        "subtypes": sorted(found),
        "complete": not incomplete,
        "reasons": sorted(incomplete),
    }


def process_xlsx(_artifact: Mapping[str, Any], stream: BinaryIO, detection: Mapping[str, Any]) -> dict[str, Any]:
    from openpyxl import load_workbook
    payload = _read_bounded(stream); warnings = _office_active_warnings(payload)
    if any("DTD or entity" in warning or "invalid OOXML" in warning for warning in warnings):
        return _result("safe-openpyxl", detection, units=[], status="unsupported", warnings=warnings, mode="tabular_parse")
    workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=False, keep_links=False)
    units = []
    try:
        for sheet in workbook.worksheets:
            units.extend(_table_units([list(row) for row in sheet.iter_rows(values_only=True)], sheet_name=sheet.title))
    finally:
        workbook.close()
    return _result("safe-openpyxl", detection, units=units, warnings=warnings, status="partial" if warnings else "ready", mode="tabular_parse")


def process_pdf(_artifact: Mapping[str, Any], stream: BinaryIO, detection: Mapping[str, Any]) -> dict[str, Any]:
    from pypdf import PdfReader
    payload = _read_bounded(stream)
    try:
        reader = PdfReader(io.BytesIO(payload), strict=False)
        if reader.is_encrypted:
            return _result("safe-pypdf", detection, units=[], status="needs_user", warnings=["encrypted PDF requires user-provided unlocked material"])
        pages = len(reader.pages)
        if pages > MAX_PDF_PAGES:
            return _result("safe-pypdf", detection, units=[], status="needs_user", warnings=["PDF exceeds 500-page processing ceiling"])
        warnings = []
        action_inspection = _inspect_pdf_actions(reader)
        if action_inspection["subtypes"]:
            warnings.append(
                "PDF action(s) reported and never executed: "
                + ", ".join(action_inspection["subtypes"])
            )
        if not action_inspection["complete"]:
            warnings.append(
                "PDF action inspection incomplete: "
                + ", ".join(action_inspection["reasons"])
            )
        units, missing = [], 0
        for page_number, page in enumerate(reader.pages, 1):
            try: text = (page.extract_text() or "").strip()
            except Exception: text = ""
            if not text: missing += 1
            units.append({"kind": "pdf_page", "content": text if text else None, "metadata": {"page": page_number, "text_available": bool(text)}, "source_span": {"kind": "pdf_region", "locator": {"page": page_number, "normalized_rect": [0.0, 0.0, 1.0, 1.0]}}, "candidate_only": True})
        if missing: warnings.append(f"text extraction unavailable on {missing} page(s); no text was fabricated")
        return _result("safe-pypdf", detection, units=units, status="partial" if missing or warnings else "ready", warnings=warnings)
    except Exception as error:
        return _result("safe-pypdf", detection, units=[], status="failed", warnings=[f"PDF format failure: {type(error).__name__}"])


def render_pdf_page_preview(payload: bytes, *, page_number: int) -> bytes | None:
    """Return a re-encoded page PNG only when an in-process renderer is available."""

    if not isinstance(page_number, int) or isinstance(page_number, bool) or page_number < 1:
        raise ValueError("PDF preview page must be positive")
    try:
        from PIL import Image
        with Image.open(io.BytesIO(payload)) as document:
            document.seek(page_number - 1)
            page = document.convert("RGB")
            output = io.BytesIO(); page.save(output, format="PNG", optimize=True)
            return output.getvalue()
    except Exception:
        return None


def process_docx(_artifact: Mapping[str, Any], stream: BinaryIO, detection: Mapping[str, Any]) -> dict[str, Any]:
    from docx import Document
    payload = _read_bounded(stream); warnings = _office_active_warnings(payload)
    if any("DTD or entity" in warning or "invalid OOXML" in warning for warning in warnings):
        return _result("safe-python-docx", detection, units=[], status="unsupported", warnings=warnings)
    try: document = Document(io.BytesIO(payload))
    except Exception as error: return _result("safe-python-docx", detection, units=[], status="unsupported" if warnings else "failed", warnings=warnings + [f"DOCX format failure: {type(error).__name__}"])
    heading, paragraph, units = "Document", 0, []
    for item in document.paragraphs:
        content = item.text.strip()
        if not content: continue
        if item.style and item.style.name.casefold().startswith("heading"):
            heading = content; paragraph = 0; continue
        units.append({"kind": "document_text", "content": content, "source_span": {"kind": "document_section", "locator": {"heading_path": [heading], "paragraph_start": paragraph, "paragraph_end": paragraph}}, "candidate_only": True}); paragraph += 1
    for table_index, table in enumerate(document.tables):
        for row_index, row in enumerate(table.rows):
            original = [cell.text for cell in row.cells]
            content = " | ".join(original)
            units.append({"kind": "document_table_row", "content": content, "provenance": {"original_value": original, "original_unit": None, "transformations": ["joined_table_cells"], "assumptions": [], "missing_data": [], "validation_status": "candidate_unvalidated"}, "source_span": {"kind": "document_section", "locator": {"heading_path": [heading, f"Table {table_index + 1}"], "paragraph_start": row_index, "paragraph_end": row_index}}, "candidate_only": True})
    return _result("safe-python-docx", detection, units=units, status="partial" if warnings else "ready", warnings=warnings)


def process_pptx(_artifact: Mapping[str, Any], stream: BinaryIO, detection: Mapping[str, Any]) -> dict[str, Any]:
    from pptx import Presentation
    payload = _read_bounded(stream); warnings = _office_active_warnings(payload)
    if any("DTD or entity" in warning or "invalid OOXML" in warning for warning in warnings):
        return _result("safe-python-pptx", detection, units=[], status="unsupported", warnings=warnings)
    try: presentation = Presentation(io.BytesIO(payload))
    except Exception as error: return _result("safe-python-pptx", detection, units=[], status="unsupported" if warnings else "failed", warnings=warnings + [f"PPTX format failure: {type(error).__name__}"])
    units = []
    for slide_number, slide in enumerate(presentation.slides, 1):
        index = 0
        for shape in slide.shapes:
            content = getattr(shape, "text", "").strip()
            if not content: continue
            units.append({"kind": "document_text", "content": content, "metadata": {"slide": slide_number, "notes_status": "not_extracted"}, "source_span": {"kind": "document_section", "locator": {"heading_path": [f"Slide {slide_number}"], "paragraph_start": index, "paragraph_end": index}}, "candidate_only": True}); index += 1
    return _result("safe-python-pptx", detection, units=units, status="partial" if warnings else "ready", warnings=warnings + ["speaker notes status reported; notes text not extracted"])


def process_image(_artifact: Mapping[str, Any], stream: BinaryIO, detection: Mapping[str, Any]) -> dict[str, Any]:
    from PIL import Image
    payload = _read_bounded(stream)
    if detection["mime"] == "image/png" and len(payload) >= 24:
        width, height = struct.unpack(">II", payload[16:24])
        if width * height > MAX_IMAGE_PIXELS:
            return _result("safe-pillow-metadata", detection, units=[], status="needs_user", warnings=["image exceeds 40 MP processing ceiling"], mode="metadata_only")
    try:
        with Image.open(io.BytesIO(payload)) as image:
            width, height = image.size; image.verify()
    except Exception as error:
        return _result("safe-pillow-metadata", detection, units=[], status="failed", warnings=[f"image format failure: {type(error).__name__}"], mode="metadata_only")
    if width * height > MAX_IMAGE_PIXELS:
        return _result("safe-pillow-metadata", detection, units=[], status="needs_user", warnings=["image exceeds 40 MP processing ceiling"], mode="metadata_only")
    unit = {"kind": "image_metadata", "metadata": {"width": width, "height": height, "visual_understanding": False}, "source_span": {"kind": "image_region", "locator": {"normalized_rect": [0.0, 0.0, 1.0, 1.0]}}, "candidate_only": True}
    return _result("safe-pillow-metadata", detection, units=[unit], status="partial", warnings=["visual understanding unavailable; metadata only"], mode="metadata_only")


def process_audio(_artifact: Mapping[str, Any], stream: BinaryIO, detection: Mapping[str, Any]) -> dict[str, Any]:
    payload = _read_bounded(stream)
    if detection["mime"] != "audio/wav":
        return _result(
            "safe-audio-metadata",
            detection,
            units=[],
            status="partial",
            warnings=[
                "duration and channel metadata unavailable for this audio container",
                "transcription is an optional unavailable candidate capability",
            ],
            mode="metadata_only",
        )
    metadata: dict[str, Any] = {"transcription_available": False}
    with wave.open(io.BytesIO(payload), "rb") as source:
        frames, rate = source.getnframes(), source.getframerate()
        duration_ms = max(1, round(frames * 1000 / rate))
        metadata.update({"duration_ms": duration_ms, "channels": source.getnchannels(), "sample_rate_hz": rate})
    unit = {"kind": "audio_metadata", "metadata": metadata, "source_span": {"kind": "audio_segment", "locator": {"start_ms": 0, "end_ms": duration_ms}}, "candidate_only": True}
    return _result("safe-audio-metadata", detection, units=[unit], status="partial", warnings=["transcription is an optional unavailable candidate capability"], mode="metadata_only")


def default_processor_registry() -> ProcessorRegistry:
    registry = ProcessorRegistry()
    registry.register("text/plain", process_text)
    registry.register("application/json", process_json)
    registry.register("application/xml", process_xml)
    registry.register("text/html", process_html)
    registry.register("text/csv", process_delimited)
    registry.register("text/tab-separated-values", process_delimited)
    registry.register("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", process_xlsx)
    registry.register("application/pdf", process_pdf)
    registry.register("application/vnd.openxmlformats-officedocument.wordprocessingml.document", process_docx)
    registry.register("application/vnd.openxmlformats-officedocument.presentationml.presentation", process_pptx)
    registry.register("image/png", process_image)
    registry.register("image/jpeg", process_image)
    registry.register("audio/wav", process_audio)
    registry.register("audio/mpeg", process_audio)
    registry.register("audio/mp4", process_audio)
    return registry
